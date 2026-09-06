from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from build_taller_bioquimica import (
    ASSETS,
    INK,
    MUTED,
    NAVY,
    ORANGE,
    OUT,
    PALE_GRAY,
    PALE_ORANGE,
    PALE_TEAL,
    STUDENT,
    TEACHER,
    TEAL,
    add_answer,
    add_callout,
    add_col_break,
    add_label_line,
    add_page_field,
    add_picture,
    add_question,
    add_response_lines,
    add_small_table,
    configure_document,
    set_columns,
    set_font,
    set_table_geometry,
    set_cell_shading,
)


NEW_PLATE = Path(r"C:\Users\andre\.codex\generated_images\01a021bf-5c05-7010-b67b-13e73df76443\exec-77ee5f74-7886-4a51-a9f4-04b55bc3f186.png")


def crop_new_assets():
    ASSETS.mkdir(parents=True, exist_ok=True)
    image = Image.open(NEW_PLATE)
    w, h = image.size
    crops = {
        "acidos_grasos_estructura.png": (0, 60, int(w * 0.30), h - 60),
        "niveles_carbohidratos.png": (int(w * 0.28), 95, int(w * 0.59), h - 80),
        "polisacaridos_lineal_ramificado.png": (int(w * 0.60), 55, w, h - 55),
    }
    for name, box in crops.items():
        image.crop(box).save(ASSETS / name)


def add_intro(doc):
    p = doc.add_paragraph(style="Title")
    p.add_run("Taller de aplicación")
    sp = doc.add_paragraph(style="Subtitle")
    sp.add_run("Estructura y función de lípidos, membranas y carbohidratos")
    add_label_line(doc, "Nombre", 42)
    info = doc.add_table(rows=1, cols=3)
    set_table_geometry(info, [3000, 3000, 3000], indent=0)
    for cell, text in zip(info.rows[0].cells, ["Grupo: __________", "Fecha: __________", "Tiempo: 60 min"]):
        set_cell_shading(cell, PALE_GRAY)
        r = cell.paragraphs[0].add_run(text)
        set_font(r, 8.5, bold=True, color=NAVY)
    add_callout(
        doc,
        "Propósito.",
        "Usar la forma, el tamaño y la organización de lípidos y carbohidratos para explicar funciones generales en células y productos farmacéuticos.",
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    r = p.add_run("Indicaciones: ")
    set_font(r, 9, bold=True, color=NAVY)
    r = p.add_run("trabaje en parejas; observe antes de responder; toda elección debe incluir una razón estructural. No se requieren rutas metabólicas ni nombres de transportadores. Puntaje total: 40 puntos.")
    set_font(r, 9, color=INK)


def build_student():
    doc = Document()
    configure_document(doc)
    add_intro(doc)

    sec = doc.add_section(WD_SECTION.CONTINUOUS)
    sec.top_margin = Inches(0.62)
    sec.bottom_margin = Inches(0.62)
    sec.left_margin = Inches(0.68)
    sec.right_margin = Inches(0.68)
    sec.header_distance = Inches(0.3)
    sec.footer_distance = Inches(0.3)
    set_columns(sec, 2, 360)

    doc.add_paragraph("ESTACIÓN 1 · UNA FRONTERA CON DOS CARAS", style="Heading 1")
    add_picture(doc, ASSETS / "membrana.png", "Figura 1. Representación simplificada de una membrana celular.")
    add_question(doc, "1a", "Señale en la figura: zona que interactúa con el agua, zona que evita el agua, una proteína y el colesterol. Escriba al lado la función general de cada componente.", 4)
    add_response_lines(doc, 4)
    add_question(doc, "1b", "Los fosfolípidos tienen una parte polar y otra apolar. Explique por qué, al colocarlos en agua, forman una bicapa y no una fila con todas las partes mezcladas.", 3)
    add_response_lines(doc, 4)
    add_question(doc, "1c", "Un detergente también posee una parte que interactúa con agua y otra con grasa. Prediga qué podría ocurrir si entra en contacto con una membrana y justifique desde su estructura.", 3)
    add_response_lines(doc, 4)

    doc.add_paragraph("ESTACIÓN 2 · RECTA O DOBLADA", style="Heading 1")
    add_picture(doc, ASSETS / "acidos_grasos_estructura.png", "Figura 2. Dos cadenas lipídicas con formas diferentes.")
    add_question(doc, "2a", "Una cadena es recta y la otra tiene una curvatura. ¿Cuál puede empaquetarse con mayor facilidad junto a otras cadenas? Represéntelo con un dibujo pequeño.", 3)
    add_response_lines(doc, 3)
    add_question(doc, "2b", "Dos membranas tienen igual cantidad de lípidos, pero una contiene más cadenas curvadas. Prediga cuál será más flexible a la misma temperatura y explique por qué.", 3)
    add_response_lines(doc, 4)

    doc.add_paragraph("ESTACIÓN 3 · EL RETO DE ATRAVESAR", style="Heading 1")
    add_callout(doc, "Tres visitantes.", "A es una molécula pequeña y apolar; B es una molécula grande con muchos grupos polares; C es un ion con carga eléctrica.", PALE_ORANGE)
    add_question(doc, "3", "Ordene A, B y C desde la que atravesaría más fácilmente la bicapa hasta la que tendría mayor dificultad. Para B y C indique qué componente general de la membrana podría ayudarles.", 5)
    add_response_lines(doc, 6)

    doc.add_paragraph("ESTACIÓN 4 · DEL BLOQUE A LA CADENA", style="Heading 1")
    add_picture(doc, ASSETS / "niveles_carbohidratos.png", "Figura 3. Una unidad, dos unidades y una cadena de unidades de carbohidrato.")
    add_question(doc, "4a", "Clasifique las tres estructuras como monosacárido, disacárido o polisacárido. Escriba el criterio que utilizó.", 3)
    add_response_lines(doc, 3)
    add_question(doc, "4b", "Una célula necesita combustible disponible con rapidez y también una reserva compacta. ¿Qué tipo de estructura elegiría para cada función? Justifique sin describir rutas metabólicas.", 3)
    add_response_lines(doc, 4)

    doc.add_paragraph("ESTACIÓN 5 · LÍNEA O RAMIFICACIÓN", style="Heading 1")
    add_picture(doc, ASSETS / "polisacaridos_lineal_ramificado.png", "Figura 4. Dos formas generales de organización de polisacáridos.")
    add_question(doc, "5a", "¿Cuál organización presenta más extremos accesibles: la lineal o la ramificada? Encierre esos extremos en la figura y explique la ventaja funcional.", 3)
    add_response_lines(doc, 3)
    add_question(doc, "5b", "Una fibra vegetal debe ser resistente; una reserva celular debe ser compacta y movilizable. Asigne la organización más adecuada a cada función y defienda su elección.", 3)
    add_response_lines(doc, 4)

    add_col_break(doc)
    doc.add_paragraph("ESTACIÓN 6 · DECISIÓN FARMACÉUTICA", style="Heading 1")
    add_callout(doc, "Problema de formulación.", "Una tableta necesita: (i) un material soluble de sabor suave que aporte volumen; (ii) un material que absorba agua y ayude a que la tableta se desarme; (iii) una cubierta vegetal resistente que no se digiera fácilmente.")
    add_small_table(doc, ["Material", "Descripción estructural general"], [
        ("Lactosa", "Disacárido pequeño y soluble"),
        ("Almidón", "Polisacárido de reserva que capta agua"),
        ("Celulosa", "Polisacárido lineal que forma fibras"),
    ], [1400, 2800])
    add_question(doc, "6", "Asigne lactosa, almidón y celulosa a las funciones i, ii y iii. En cada caso, conecte una característica estructural con la función elegida.", 4)
    add_response_lines(doc, 6)

    doc.add_paragraph("CIERRE · EXPLIQUE SIN MEMORIZAR", style="Heading 1")
    add_question(doc, "7", "En 60–80 palabras explique por qué dos biomoléculas formadas por elementos semejantes pueden cumplir funciones distintas. Use un ejemplo de lípidos o carbohidratos e incluya forma, organización y función.", 3)
    add_response_lines(doc, 7)
    add_callout(doc, "Antes de entregar.", "Compruebe que cada respuesta mencione una característica observable de la estructura y la relacione con una función.", PALE_GRAY)

    doc.core_properties.title = "Taller: estructura y función de lípidos, membranas y carbohidratos"
    doc.core_properties.subject = "Bioquímica para Química Farmacéutica"
    doc.core_properties.author = "Docencia de Bioquímica"
    doc.save(STUDENT)


def build_teacher():
    doc = Document()
    configure_document(doc, teacher=True)
    doc.add_paragraph("Guía docente", style="Title")
    doc.add_paragraph("Solucionario orientativo · Estructura y función de lípidos, membranas y carbohidratos", style="Subtitle")
    add_callout(doc, "Alcance.", "Las respuestas deben apoyarse en forma, tamaño, polaridad y organización general. No se exige metabolismo, cinética, enlaces glucosídicos específicos ni clasificación detallada del transporte.")
    doc.add_paragraph("Clave de respuestas", style="Heading 1")
    add_answer(doc, "1a (4 pt)", "Las cabezas de los fosfolípidos corresponden a la zona que interactúa con agua; las colas forman la región que evita el agua y actúa como barrera. La proteína participa en comunicación o paso selectivo de sustancias. El colesterol se ubica entre las colas y contribuye a regular flexibilidad y organización.")
    add_answer(doc, "1b–1c (6 pt)", "En agua, las partes polares quedan expuestas y las apolares se agrupan alejadas del agua, originando la bicapa. Un detergente puede introducir su región apolar entre las colas y mantener su parte polar hacia el agua; así desorganiza o fragmenta la membrana.")
    add_answer(doc, "2a–2b (6 pt)", "La cadena recta se empaqueta con mayor facilidad. La curvatura deja espacios e impide el contacto estrecho; por ello una membrana con más cadenas curvadas suele ser más flexible a igual temperatura.")
    add_answer(doc, "3 (5 pt)", "Orden esperado: A, luego B y finalmente C. La molécula pequeña y apolar es compatible con el interior apolar. B y C encuentran una barrera desfavorable por su polaridad o carga; proteínas de membrana pueden facilitar su paso.")
    add_answer(doc, "4a–4b (6 pt)", "Una unidad = monosacárido; dos = disacárido; muchas = polisacárido. Para disponibilidad rápida se elige una estructura pequeña, mientras que una cadena larga permite almacenar muchas unidades en una sola estructura. Basta una justificación general coherente.")
    add_answer(doc, "5a–5b (6 pt)", "La estructura ramificada presenta más extremos accesibles y permite acceso desde varios puntos; además puede organizarse de manera compacta, apropiada para reserva. Las cadenas lineales paralelas pueden asociarse y formar fibras resistentes, apropiadas para soporte estructural.")
    add_answer(doc, "6 (4 pt)", "i = lactosa, por ser pequeña y soluble; ii = almidón, porque el polisacárido capta agua y favorece el desarme; iii = celulosa, porque sus cadenas lineales forman fibras resistentes y no se digieren con facilidad en humanos.")
    add_answer(doc, "7 (3 pt)", "Debe mostrar que la función depende de cómo se ordenan las unidades, no solo de los elementos químicos presentes. Son válidos ejemplos como cadenas lipídicas rectas o curvadas y polisacáridos lineales o ramificados.")
    doc.add_page_break()
    doc.add_paragraph("Rúbrica transversal", style="Heading 1")
    add_small_table(doc, ["Nivel", "Evidencia en la respuesta"], [
        ("Logrado", "Identifica la estructura y explica cómo permite la función."),
        ("En proceso", "Identifica estructura o función, pero la relación queda incompleta."),
        ("Inicial", "Enumera términos sin usar la imagen ni establecer una relación."),
    ], [1500, 7860])
    doc.core_properties.title = "Guía docente: estructura y función"
    doc.core_properties.author = "Docencia de Bioquímica"
    doc.save(TEACHER)


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    crop_new_assets()
    build_student()
    build_teacher()
    print(STUDENT)
    print(TEACHER)
